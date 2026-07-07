from __future__ import annotations

import argparse
import configparser
import csv
import json
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def get_app_dir() -> Path:
	if getattr(sys, "frozen", False):
		return Path(sys.executable).resolve().parent
	return Path(__file__).resolve().parent


DEFAULT_CONFIG_FILE = get_app_dir() / "get_log_data.ini"
REQUIRED_COLUMNS = {"uid", "timestamp", "activity", "attribute"}
CLONE_ACTIVITY = "Result_Clone"
SUCCESS_RESULT = "success"


def load_export_config(config_path: Path) -> dict[str, str]:
	config = configparser.RawConfigParser()
	config.read(config_path, encoding="utf-8")
	if not config.has_section("export"):
		raise RuntimeError(f"配置文件缺少 [export] 节: {config_path}")
	return dict(config.items("export"))


def resolve_input_path(export_config: dict[str, str]) -> Path:
	output = export_config.get("output")
	if not output:
		raise RuntimeError("配置文件 [export] 缺少 output")

	output_path = Path(output)
	if output_path.is_absolute():
		return output_path
	return Path(export_config.get("output_dir", ".")) / output_path


def default_report_path(input_path: Path) -> Path:
	return input_path.with_name(f"{input_path.stem}_clone_report.xlsx")


def parse_attribute(value: str) -> dict[str, Any]:
	value = (value or "").strip()
	if not value:
		return {"result": "", "parse_error": "empty attribute"}
	try:
		parsed = json.loads(value)
	except json.JSONDecodeError as exc:
		return {"result": "", "parse_error": str(exc)}
	if not isinstance(parsed, dict):
		return {"result": "", "parse_error": "attribute is not a JSON object"}
	return parsed


def get_clone_result(attribute: dict[str, Any]) -> tuple[str, str, str]:
    result = str(attribute.get("result") or "").strip()
    message = str(attribute.get("message") or "").strip()
    time_used = str(attribute.get("time_used") or "").strip()
    if result:
        return result, message, time_used

    clone_result = attribute.get(CLONE_ACTIVITY)
    if not isinstance(clone_result, dict):
        return "", message, time_used

    disk_statuses: list[str] = []
    disk_messages: list[str] = []
    disk_time_used: list[str] = []
    for disk_name, disk_result in sorted(clone_result.items()):
        if not isinstance(disk_result, dict):
            continue
        disk_status = str(disk_result.get("result") or "").strip()
        disk_message = str(disk_result.get("message") or "").strip()
        disk_time = str(disk_result.get("time_used") or "").strip()
        if disk_status:
            disk_statuses.append(disk_status)
            disk_messages.append(f"{disk_name}:{disk_status}")
        if disk_message:
            disk_messages.append(f"{disk_name}:{disk_message}")
        if disk_time:
            disk_time_used.append(f"{disk_name}:{disk_time}")

    result = SUCCESS_RESULT if disk_statuses and all(status == SUCCESS_RESULT for status in disk_statuses) else "failure"
    return result, "; ".join(disk_messages), "; ".join(disk_time_used)


def timestamp_sort_value(row: dict[str, str]) -> int:
	try:
		return int((row.get("timestamp") or "").strip())
	except ValueError:
		return 0


def format_timestamp(row: dict[str, str]) -> str:
	timestamp = timestamp_sort_value(row)
	if not timestamp:
		return ""
	return datetime.fromtimestamp(timestamp).strftime("%Y-%m-%d %H:%M:%S")


def format_rate(count: int, total: int) -> float:
	return count / total if total else 0


def is_failed_clone_record(item: dict[str, Any]) -> bool:
	activity = (item["row"].get("activity") or "").strip()
	return activity == CLONE_ACTIVITY and item["result"] != SUCCESS_RESULT


def is_success_clone_record(item: dict[str, Any]) -> bool:
	activity = (item["row"].get("activity") or "").strip()
	return activity == CLONE_ACTIVITY and item["result"] == SUCCESS_RESULT


def read_rows(input_path: Path) -> tuple[list[str], list[dict[str, str]]]:
	with input_path.open("r", encoding="utf-8-sig", newline="") as f:
		reader = csv.DictReader(f)
		fieldnames = list(reader.fieldnames or [])
		missing_columns = REQUIRED_COLUMNS - set(fieldnames)
		if missing_columns:
			raise RuntimeError(f"CSV 缺少必要列: {', '.join(sorted(missing_columns))}")
		return fieldnames, list(reader)


def parse_row(row: dict[str, str]) -> dict[str, Any]:
	uid = (row.get("uid") or "").strip()
	attribute = parse_attribute(row.get("attribute") or "")
	result, message, time_used = get_clone_result(attribute)
	parse_error = str(attribute.get("parse_error") or "").strip()
	return {
		"row": row,
		"uid": uid,
		"result": result or "parse_error",
		"message": message,
		"time_used": time_used,
		"parse_error": parse_error,
	}


def collect_report_data(rows: list[dict[str, str]]) -> dict[str, Any]:
	parsed_rows = [parse_row(row) for row in rows]
	clone_rows = [item for item in parsed_rows if (item["row"].get("activity") or "").strip() == CLONE_ACTIVITY]
	uid_results: dict[str, set[str]] = {}
	result_counts: dict[str, int] = {}

	for item in clone_rows:
		uid = item["uid"]
		result_key = item["result"]
		result_counts[result_key] = result_counts.get(result_key, 0) + 1
		if uid:
			uid_results.setdefault(uid, set()).add(result_key)

	total_uids = len(uid_results)
	success_uids = {uid for uid, results in uid_results.items() if SUCCESS_RESULT in results}
	failed_uids = {uid for uid, results in uid_results.items() if any(result != SUCCESS_RESULT for result in results)}
	no_failed_uids = set(uid_results) - failed_uids
	failed_uid_records = [item for item in parsed_rows if item["uid"] in failed_uids]
	failed_uid_records.sort(key=lambda item: (item["uid"], timestamp_sort_value(item["row"])))

	return {
		"input_records": len(rows),
		"total_records": len(clone_rows),
		"total_uids": total_uids,
		"success_uids": success_uids,
		"failed_uids": failed_uids,
		"no_failed_uids": no_failed_uids,
		"result_counts": result_counts,
		"failed_uid_records": failed_uid_records,
	}


def append_rows(ws, headers: list[str], rows: list[list[object]]) -> None:
	ws.append(headers)
	for row in rows:
		ws.append(row)


def build_failed_headers(fieldnames: list[str]) -> list[str]:
	headers: list[str] = []
	for field in fieldnames:
		headers.append(field)
		if field == "timestamp":
			headers.append("timestamp_time")
	return headers + ["parsed_result", "parsed_message", "parsed_time_used", "parse_error"]


def build_failed_row(fieldnames: list[str], item: dict[str, Any]) -> list[object]:
	row = item["row"]
	values: list[object] = []
	for field in fieldnames:
		values.append(row.get(field, ""))
		if field == "timestamp":
			values.append(format_timestamp(row))
	return values + [item["result"], item["message"], item["time_used"], item["parse_error"]]


def style_sheet(ws, freeze_panes: str = "A2") -> None:
	header_fill = PatternFill("solid", fgColor="D9EAF7")
	for cell in ws[1]:
		cell.font = Font(bold=True)
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center")

	ws.freeze_panes = freeze_panes
	ws.auto_filter.ref = ws.dimensions
	for column_index, column_cells in enumerate(ws.columns, start=1):
		max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
		ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 60)


def write_xlsx(fieldnames: list[str], report_data: dict[str, Any], output_path: Path) -> None:
	output_path.parent.mkdir(parents=True, exist_ok=True)
	wb = Workbook()

	summary_ws = wb.active
	summary_ws.title = "summary"
	total_uids = report_data["total_uids"]
	success_count = len(report_data["success_uids"])
	failed_count = len(report_data["failed_uids"])
	no_failed_count = len(report_data["no_failed_uids"])
	append_rows(
		summary_ws,
		["metric", "value", "rate"],
		[
			["input_records", report_data["input_records"], None],
			["total_records", report_data["total_records"], None],
			["total_uids", total_uids, None],
			["success_uids", success_count, format_rate(success_count, total_uids)],
			["failed_uids", failed_count, format_rate(failed_count, total_uids)],
			["no_failed_uids", no_failed_count, format_rate(no_failed_count, total_uids)],
		],
	)
	for row_index in range(2, summary_ws.max_row + 1):
		summary_ws.cell(row=row_index, column=3).number_format = "0.00%"
	style_sheet(summary_ws)

	result_ws = wb.create_sheet("result_counts")
	result_rows = [[result, count] for result, count in sorted(report_data["result_counts"].items())]
	append_rows(result_ws, ["result", "records"], result_rows)
	style_sheet(result_ws)

	failed_ws = wb.create_sheet("failed_uid_records")
	failed_headers = build_failed_headers(fieldnames)
	failed_rows = [build_failed_row(fieldnames, item) for item in report_data["failed_uid_records"]]
	append_rows(failed_ws, failed_headers, failed_rows)
	red_font = Font(color="FFFF0000")
	green_font = Font(color="FF008000")
	for row_index, item in enumerate(report_data["failed_uid_records"], start=2):
		if is_success_clone_record(item):
			font = green_font
		elif is_failed_clone_record(item):
			font = red_font
		else:
			continue
		for cell in failed_ws[row_index]:
			cell.font = font
	style_sheet(failed_ws)

	wb.save(output_path)


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description="统计克隆结果 CSV，并生成 xlsx 报告")
	parser.add_argument("config_file", nargs="?", default=str(DEFAULT_CONFIG_FILE), help="get_log_data.ini 路径")
	parser.add_argument("--output", help="xlsx 输出路径，默认与 CSV 同目录，文件名追加 _clone_report.xlsx")
	return parser.parse_args()


def main() -> None:
	args = parse_args()
	config_path = Path(args.config_file)
	export_config = load_export_config(config_path)
	input_path = resolve_input_path(export_config)
	output_path = Path(args.output) if args.output else default_report_path(input_path)

	fieldnames, rows = read_rows(input_path)
	report_data = collect_report_data(rows)
	write_xlsx(fieldnames, report_data, output_path)

	print(f"Input: {input_path}")
	print(f"Input Rows: {report_data['input_records']}")
	print(f"Rows: {report_data['total_records']}")
	print(f"UIDs: {report_data['total_uids']}")
	print(f"Success UIDs: {len(report_data['success_uids'])}")
	print(f"Failed UIDs: {len(report_data['failed_uids'])}")
	print(f"No Failed UIDs: {len(report_data['no_failed_uids'])}")
	print(f"Failed UID Records: {len(report_data['failed_uid_records'])}")
	print(f"Output: {output_path}")


if __name__ == "__main__":
	main()

