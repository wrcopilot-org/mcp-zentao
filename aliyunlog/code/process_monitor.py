from __future__ import annotations

import argparse
import signal
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import psutil
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def get_app_dir() -> Path:
	if getattr(sys, "frozen", False):
		return Path(sys.executable).resolve().parent
	return Path(__file__).resolve().parent


DEFAULT_OUTPUT_FILE = get_app_dir().parent / "output" / "process_monitor.xlsx"
METRIC_NAMES = ["threads", "handles", "memory_mb", "cpu_time_seconds", "run_time_seconds"]


@dataclass
class ProcessSample:
	timestamp: datetime
	pid: int
	name: str
	exe: str
	cmdline: str
	username: str
	status: str
	threads: int
	handles: int | None
	memory_mb: float
	cpu_time_seconds: float
	run_time_seconds: float


@dataclass
class MetricStats:
	minimum: float | None = None
	maximum: float | None = None
	latest: float | None = None

	def update(self, value: float | int | None) -> None:
		if value is None:
			return
		value = float(value)
		self.minimum = value if self.minimum is None else min(self.minimum, value)
		self.maximum = value if self.maximum is None else max(self.maximum, value)
		self.latest = value


@dataclass
class ProcessStats:
	pid: int
	name: str
	exe: str
	cmdline: str
	username: str
	first_seen: datetime
	last_seen: datetime
	sample_count: int = 0
	metrics: dict[str, MetricStats] = field(default_factory=lambda: {name: MetricStats() for name in METRIC_NAMES})

	def update(self, sample: ProcessSample) -> None:
		self.name = sample.name
		self.exe = sample.exe
		self.cmdline = sample.cmdline
		self.username = sample.username
		self.last_seen = sample.timestamp
		self.sample_count += 1
		self.metrics["threads"].update(sample.threads)
		self.metrics["handles"].update(sample.handles)
		self.metrics["memory_mb"].update(sample.memory_mb)
		self.metrics["cpu_time_seconds"].update(sample.cpu_time_seconds)
		self.metrics["run_time_seconds"].update(sample.run_time_seconds)


def current_process_ids() -> set[int]:
	return {process.info["pid"] for process in psutil.process_iter(["pid"])}


def safe_join_cmdline(cmdline: list[str] | None) -> str:
	return " ".join(cmdline or [])


def collect_process_sample(process: psutil.Process, timestamp: datetime) -> ProcessSample | None:
	try:
		with process.oneshot():
			cpu_times = process.cpu_times()
			create_time = process.create_time()
			memory_info = process.memory_info()
			handles = process.num_handles() if hasattr(process, "num_handles") else None
			return ProcessSample(
				timestamp=timestamp,
				pid=process.pid,
				name=process.name(),
				exe=process.exe(),
				cmdline=safe_join_cmdline(process.cmdline()),
				username=process.username(),
				status=process.status(),
				threads=process.num_threads(),
				handles=handles,
				memory_mb=round(memory_info.rss / 1024 / 1024, 2),
				cpu_time_seconds=round(cpu_times.user + cpu_times.system, 2),
				run_time_seconds=round(max(0.0, time.time() - create_time), 2),
			)
	except (psutil.AccessDenied, psutil.NoSuchProcess, psutil.ZombieProcess):
		return None


def collect_new_processes(known_pids: set[int]) -> list[ProcessSample]:
	timestamp = datetime.now()
	samples: list[ProcessSample] = []
	for process in psutil.process_iter(["pid"]):
		pid = process.info["pid"]
		if pid in known_pids:
			continue
		sample = collect_process_sample(process, timestamp)
		known_pids.add(pid)
		if sample is not None:
			samples.append(sample)
	return samples


def collect_tracked_processes(tracked_pids: set[int]) -> list[ProcessSample]:
	timestamp = datetime.now()
	samples: list[ProcessSample] = []
	for pid in list(tracked_pids):
		try:
			process = psutil.Process(pid)
		except psutil.NoSuchProcess:
			tracked_pids.discard(pid)
			continue
		sample = collect_process_sample(process, timestamp)
		if sample is None:
			tracked_pids.discard(pid)
		else:
			samples.append(sample)
	return samples


def update_stats(stats_by_pid: dict[int, ProcessStats], samples: list[ProcessSample]) -> None:
	for sample in samples:
		stats = stats_by_pid.get(sample.pid)
		if stats is None:
			stats = ProcessStats(
				pid=sample.pid,
				name=sample.name,
				exe=sample.exe,
				cmdline=sample.cmdline,
				username=sample.username,
				first_seen=sample.timestamp,
				last_seen=sample.timestamp,
			)
			stats_by_pid[sample.pid] = stats
		stats.update(sample)


def format_datetime(value: datetime) -> str:
	return value.strftime("%Y-%m-%d %H:%M:%S")


def style_sheet(ws, freeze_panes: str = "A2") -> None:
	header_fill = PatternFill("solid", fgColor="D9EAF7")
	for cell in ws[1]:
		cell.font = Font(bold=True)
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center")

	ws.freeze_panes = freeze_panes
	ws.auto_filter.ref = ws.dimensions
	for column_index, column_cells in enumerate(ws.columns, start=1):
		max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
		ws.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 12), 80)


def write_report(stats_by_pid: dict[int, ProcessStats], samples: list[ProcessSample], output_path: Path) -> None:
	output_path.parent.mkdir(parents=True, exist_ok=True)
	wb = Workbook()
	summary_ws = wb.active
	summary_ws.title = "summary"

	summary_headers = ["pid", "name", "exe", "cmdline", "username", "first_seen", "last_seen", "sample_count"]
	for metric in METRIC_NAMES:
		summary_headers.extend([f"{metric}_min", f"{metric}_max", f"{metric}_latest"])
	summary_ws.append(summary_headers)

	for stats in sorted(stats_by_pid.values(), key=lambda item: (item.first_seen, item.pid)):
		row: list[object] = [
			stats.pid,
			stats.name,
			stats.exe,
			stats.cmdline,
			stats.username,
			format_datetime(stats.first_seen),
			format_datetime(stats.last_seen),
			stats.sample_count,
		]
		for metric in METRIC_NAMES:
			metric_stats = stats.metrics[metric]
			row.extend([metric_stats.minimum, metric_stats.maximum, metric_stats.latest])
		summary_ws.append(row)
	style_sheet(summary_ws, "I2")

	samples_ws = wb.create_sheet("samples")
	samples_ws.append(
		[
			"timestamp",
			"pid",
			"name",
			"exe",
			"cmdline",
			"username",
			"status",
			"threads",
			"handles",
			"memory_mb",
			"cpu_time_seconds",
			"run_time_seconds",
		]
	)
	for sample in samples:
		samples_ws.append(
			[
				format_datetime(sample.timestamp),
				sample.pid,
				sample.name,
				sample.exe,
				sample.cmdline,
				sample.username,
				sample.status,
				sample.threads,
				sample.handles,
				sample.memory_mb,
				sample.cpu_time_seconds,
				sample.run_time_seconds,
			]
		)
	style_sheet(samples_ws, "H2")

	wb.save(output_path)


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description="监控新增进程，采集线程/句柄/内存/CPU 时间/运行时间，并生成 xlsx")
	parser.add_argument("--interval", type=float, default=2.0, help="扫描间隔，单位秒，默认 2 秒")
	parser.add_argument("--duration", type=float, default=0.0, help="运行时长，单位秒；0 表示一直运行到 Ctrl+C")
	parser.add_argument("--output", default=str(DEFAULT_OUTPUT_FILE), help="xlsx 输出路径")
	parser.add_argument("--include-existing", action="store_true", help="启动时把已经存在的进程也纳入监控")
	return parser.parse_args()


def main() -> None:
	args = parse_args()
	output_path = Path(args.output)
	known_pids = set() if args.include_existing else current_process_ids()
	tracked_pids: set[int] = set()
	stats_by_pid: dict[int, ProcessStats] = {}
	all_samples: list[ProcessSample] = []
	start_time = time.monotonic()
	stop_requested = False

	def request_stop(signum, frame) -> None:
		nonlocal stop_requested
		stop_requested = True

	signal.signal(signal.SIGINT, request_stop)

	print(f"Process monitor started. Output: {output_path}")
	print("Press Ctrl+C to stop and write report.")
	try:
		while not stop_requested:
			tracked_samples = collect_tracked_processes(tracked_pids)
			new_samples = collect_new_processes(known_pids)
			for sample in new_samples:
				tracked_pids.add(sample.pid)

			iteration_samples = tracked_samples + new_samples
			all_samples.extend(iteration_samples)
			update_stats(stats_by_pid, iteration_samples)

			if new_samples:
				print(f"{format_datetime(datetime.now())} found {len(new_samples)} new process(es)")

			if args.duration and time.monotonic() - start_time >= args.duration:
				break
			time.sleep(max(args.interval, 0.2))
	finally:
		write_report(stats_by_pid, all_samples, output_path)
		print(f"Processes: {len(stats_by_pid)}")
		print(f"Samples: {len(all_samples)}")
		print(f"Output: {output_path}")


if __name__ == "__main__":
	main()