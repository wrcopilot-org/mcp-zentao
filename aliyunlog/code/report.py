from __future__ import annotations

import argparse
import configparser
import csv
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def get_app_dir() -> Path:
	if getattr(sys, "frozen", False):
		return Path(sys.executable).resolve().parent
	return Path(__file__).resolve().parent


DEFAULT_CONFIG_FILE = get_app_dir() / "get_log_data.ini"
REQUIRED_COLUMNS = {"vid", "uid", "window", "activity"}
VARIANT_ORDER = {"A": 0, "B": 1}
CHANNEL_ORDER = {"free": 0, "trial": 1}

RateSpec = tuple[int, int, int]


def load_export_config(config_path: Path) -> dict[str, str]:
	config = configparser.RawConfigParser()
	config.read(config_path, encoding="utf-8")
	if not config.has_section("export"):
		raise RuntimeError(f"配置文件缺少 [export] 节: {config_path}")
	return dict(config.items("export"))


def resolve_input_path(export_config: dict[str, str]) -> Path:
	output = export_config.get("output")
	if not output:
		raise RuntimeError("配置文件 [export] 缺少 output")

	output_path = Path(output)
	if output_path.is_absolute():
		return output_path
	return Path(export_config.get("output_dir", ".")) / output_path


def default_report_path(input_path: Path) -> Path:
	return input_path.with_name(f"{input_path.stem}_report.xlsx")


def short_vid_name(vid: str) -> str:
	return vid.removeprefix("epm2030_")


def parse_vid_name(vid: str) -> tuple[str, str, str] | None:
	parts = short_vid_name(vid).split("_")
	if len(parts) < 3 or parts[1] not in VARIANT_ORDER:
		return None
	version = parts[2][1:] if parts[2][:1].isalpha() else parts[2]
	return parts[0], parts[1], version


def vid_sort_key(vid: str) -> tuple[int, str, int, str, str]:
	parsed = parse_vid_name(vid)
	if not parsed:
		return (99, short_vid_name(vid), 99, "", "")
	channel, variant, version = parsed
	version_number = int(version) if version.isdigit() else 0
	return (CHANNEL_ORDER.get(channel, 50), channel, version_number, variant, short_vid_name(vid))


def build_vid_column_groups(vids: list[str]) -> list[list[str]]:
	parsed_groups: dict[tuple[str, str], dict[str, str]] = defaultdict(dict)
	unparsed_vids: list[str] = []
	for vid in vids:
		parsed = parse_vid_name(vid)
		if not parsed:
			unparsed_vids.append(vid)
			continue
		channel, variant, version = parsed
		parsed_groups[(channel, version)][variant] = vid

	groups: list[list[str]] = []
	for key in sorted(
		parsed_groups,
		key=lambda item: (
			CHANNEL_ORDER.get(item[0], 50),
			item[0],
			int(item[1]) if item[1].isdigit() else 0,
			item[1],
		),
	):
		variants = parsed_groups[key]
		ordered_vids = [variants[variant] for variant in sorted(variants, key=lambda value: VARIANT_ORDER[value])]
		groups.append(ordered_vids)

	groups.extend([[vid] for vid in sorted(unparsed_vids, key=short_vid_name)])
	return groups


def collect_stats(input_path: Path) -> tuple[list[str], list[list[object]], list[RateSpec]]:
	group_stats: dict[tuple[str, str], dict[str, dict[str, object]]] = defaultdict(
		lambda: defaultdict(lambda: {"uids": set(), "all": 0})
	)
	vids: set[str] = set()

	with input_path.open("r", encoding="utf-8-sig", newline="") as f:
		reader = csv.DictReader(f)
		missing_columns = REQUIRED_COLUMNS - set(reader.fieldnames or [])
		if missing_columns:
			raise RuntimeError(f"CSV 缺少必要列: {', '.join(sorted(missing_columns))}")

		for row in reader:
			vid = (row.get("vid") or "").strip()
			window = (row.get("window") or "").strip()
			activity = (row.get("activity") or "").strip()
			uid = (row.get("uid") or "").strip()
			if not vid:
				continue

			vids.add(vid)
			stats = group_stats[(window, activity)][vid]
			stats["all"] = int(stats["all"]) + 1
			if uid:
				stats["uids"].add(uid)

	vid_groups = build_vid_column_groups(sorted(vids, key=vid_sort_key))
	headers = ["window", "activity"]
	rate_specs: list[RateSpec] = []
	column_index = len(headers) + 1
	for vid_group in vid_groups:
		for vid in vid_group:
			short_name = short_vid_name(vid)
			headers.extend([f"{short_name}-uid", f"{short_name}-all"])
			column_index += 2
		if len(vid_group) == 2:
			a_uid_column = column_index - 4
			a_all_column = column_index - 3
			b_uid_column = column_index - 2
			b_all_column = column_index - 1
			headers.extend(["uid-rate", "all-rate"])
			rate_specs.extend(
				[
					(column_index, a_uid_column, b_uid_column),
					(column_index + 1, a_all_column, b_all_column),
				]
			)
			column_index += 2

	rows: list[list[object]] = []
	for window, activity in sorted(group_stats):
		row: list[object] = [window, activity]
		for vid_group in vid_groups:
			for vid in vid_group:
				stats = group_stats[(window, activity)].get(vid, {"uids": set(), "all": 0})
				row.extend([len(stats["uids"]), int(stats["all"])])
			if len(vid_group) == 2:
				row.extend([None, None])
		rows.append(row)
	return headers, rows, rate_specs


def write_xlsx(headers: list[str], rows: list[list[object]], rate_specs: list[RateSpec], output_path: Path) -> None:
	output_path.parent.mkdir(parents=True, exist_ok=True)
	wb = Workbook()
	ws = wb.active
	ws.title = "report"

	ws.append(headers)
	for row in rows:
		ws.append(row)

	for row_index in range(2, ws.max_row + 1):
		for rate_column, base_column, compare_column in rate_specs:
			base_cell = f"{get_column_letter(base_column)}{row_index}"
			compare_cell = f"{get_column_letter(compare_column)}{row_index}"
			cell = ws.cell(row=row_index, column=rate_column)
			cell.value = f"=({compare_cell}-{base_cell})/{base_cell}"
			cell.number_format = "0.00%"

	header_fill = PatternFill("solid", fgColor="D9EAF7")
	for cell in ws[1]:
		cell.font = Font(bold=True)
		cell.fill = header_fill
		cell.alignment = Alignment(horizontal="center")

	ws.freeze_panes = "C2"
	ws.auto_filter.ref = ws.dimensions
	for column_index, header in enumerate(headers, start=1):
		width = min(max(len(str(header)) + 2, 12), 42)
		ws.column_dimensions[get_column_letter(column_index)].width = width

	wb.save(output_path)


def parse_args() -> argparse.Namespace:
	parser = argparse.ArgumentParser(description="按 vid/window/activity 统计 get_log_data 导出的 CSV，并生成 xlsx")
	parser.add_argument("config_file", nargs="?", default=str(DEFAULT_CONFIG_FILE), help="get_log_data.ini 路径")
	parser.add_argument("--output", help="xlsx 输出路径，默认与 CSV 同目录，文件名追加 _report.xlsx")
	return parser.parse_args()


def main() -> None:
	args = parse_args()
	config_path = Path(args.config_file)
	export_config = load_export_config(config_path)
	input_path = resolve_input_path(export_config)
	output_path = Path(args.output) if args.output else default_report_path(input_path)

	headers, rows, rate_specs = collect_stats(input_path)
	write_xlsx(headers, rows, rate_specs, output_path)

	print(f"Input: {input_path}")
	print(f"Rows: {len(rows)}")
	print(f"Columns: {len(headers)}")
	print(f"Output: {output_path}")


if __name__ == "__main__":
	main()
